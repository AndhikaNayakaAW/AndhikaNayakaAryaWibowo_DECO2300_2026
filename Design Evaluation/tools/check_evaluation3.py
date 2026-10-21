"""Check the Evaluation 3 submission artifacts and their evidence boundaries.

Run with Python containing python-docx and PyMuPDF.
"""

from pathlib import Path
import re
import unittest
from zipfile import ZipFile

from docx import Document
import pymupdf


ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT / 'Design Evaluation'
STEM = 'Design Evaluation 3 and Reflection'


class Evaluation3Checks(unittest.TestCase):
    def test_docx_is_valid_and_a4(self):
        path = FOLDER / f'{STEM}.docx'
        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
        document = Document(path)
        for section in document.sections:
            self.assertAlmostEqual(section.page_width.mm, 210, delta=0.1)
            self.assertAlmostEqual(section.page_height.mm, 297, delta=0.1)
            for margin in (section.top_margin, section.bottom_margin,
                           section.left_margin, section.right_margin):
                self.assertAlmostEqual(margin.inches, 1, places=2)
            self.assertFalse(''.join(p.text for p in section.footer.paragraphs))
            self.assertIn('PAGE', section.header._element.xml)

    def test_body_typography(self):
        document = Document(FOLDER / f'{STEM}.docx')
        for paragraph in document.paragraphs:
            if not paragraph.text.strip():
                continue
            self.assertEqual(paragraph.paragraph_format.line_spacing, 2)
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                self.assertEqual(run.font.name, 'Times New Roman')
                self.assertEqual(run.font.size.pt, 12)
                self.assertEqual(str(run.font.color.rgb), '000000')

    def test_assessed_page_budget(self):
        with pymupdf.open(FOLDER / f'{STEM}.pdf') as pdf:
            pages = [p.get_text() for p in pdf]
        self.assertIn('Objective and Validation Metrics', pages[1])
        self.assertIn('Analysis and Insights', pages[2])
        self.assertIn('Evaluation of Aims', pages[3])
        self.assertIn('Final Reflection', pages[4])
        self.assertIn('Appendix A', pages[5])
        for number, text in enumerate(pages, 1):
            self.assertEqual(text.split()[0], str(number))
            self.assertGreater(len(text.split()), 8, f'Empty page {number}')

    def test_pdf_uses_black_times_and_a4(self):
        with pymupdf.open(FOLDER / f'{STEM}.pdf') as pdf:
            for page in pdf:
                self.assertAlmostEqual(page.rect.width, 595.28, delta=0.2)
                self.assertAlmostEqual(page.rect.height, 841.89, delta=0.2)
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines', []):
                        for span in line.get('spans', []):
                            if span['text'].strip():
                                self.assertIn('TimesNewRoman', span['font'])
                                self.assertEqual(span['color'], 0)

    def test_evidence_limits_and_reflection_are_present(self):
        with pymupdf.open(FOLDER / f'{STEM}.pdf') as pdf:
            text = '\n'.join(p.get_text() for p in pdf)
        compact = re.sub(r'\s+', ' ', text)
        for phrase in ('No separate Week 12 dataset', 'notes only', '15-second',
                       '90%', 'partially confirmed', 'passthrough AR',
                       'Prototype Session Review', 'Methodology and Concept',
                       'Improvements and Extensions'):
            self.assertIn(phrase, compact)
        for symbol in (';', '—', '–', '--'):
            self.assertNotIn(symbol, text)

    def test_recording_links_are_preserved(self):
        with pymupdf.open(FOLDER / f'{STEM}.pdf') as pdf:
            urls = {link.get('uri') for p in pdf for link in p.get_links()}
        for evidence in ('1kFrLP1ncaq4-GqdgPC0OP1HqeC7Bd_w9',
                         '1gNw8L9R59p3xLFg-9AMe9GkxqopXuk_R'):
            self.assertTrue(any(evidence in (url or '') for url in urls))

    def test_readme_and_wiki_navigation(self):
        readme = (ROOT / 'README.md').read_text()
        self.assertIn('Design%20Evaluation%203%20and%20Reflection.pdf', readme)
        for name in ('Home.md', 'Design-Evaluation-3-and-Reflection.md'):
            self.assertTrue((ROOT / 'wiki' / name).is_file())
        self.assertIn('partially confirmed',
                      (ROOT / 'wiki' / 'Design-Evaluation-3-and-Reflection.md').read_text())


if __name__ == '__main__':
    unittest.main(verbosity=2)
