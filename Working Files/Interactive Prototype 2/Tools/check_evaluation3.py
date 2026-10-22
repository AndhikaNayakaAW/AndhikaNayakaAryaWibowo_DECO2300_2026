"""Check the Evaluation 3 submission artifacts and their evidence boundaries.

Run with Python containing python-docx and PyMuPDF.
"""

from pathlib import Path
import re
import unittest
from zipfile import ZipFile

from docx import Document
import pymupdf


ROOT = Path(__file__).resolve().parents[3]
FOLDER = ROOT / 'Working Files' / 'Interactive Prototype 2' / 'Evaluation'
PDF_FOLDER = ROOT / 'Interactive Prototype 2' / 'Evaluation'
STEM = 'Design Evaluation 3 and Reflection'


class Evaluation3Checks(unittest.TestCase):
    def test_docx_is_valid_and_a4(self):
        path = FOLDER / f'{STEM}.docx'
        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
        document = Document(path)
        for section in document.sections:
            self.assertAlmostEqual(section.page_width.mm, 210, delta=0.1)
            self.assertAlmostEqual(section.page_height.mm, 297, delta=0.1)
            for margin in (section.top_margin, section.bottom_margin,
                           section.left_margin, section.right_margin):
                self.assertAlmostEqual(margin.inches, 1, places=2)
            self.assertFalse(''.join(p.text for p in section.footer.paragraphs))
            self.assertIn('PAGE', section.header._element.xml)

    def test_body_typography(self):
        document = Document(FOLDER / f'{STEM}.docx')
        for paragraph in document.paragraphs:
            if not paragraph.text.strip():
                continue
            self.assertEqual(paragraph.paragraph_format.line_spacing, 2)
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                self.assertEqual(run.font.name, 'Times New Roman')
                self.assertEqual(run.font.size.pt, 12)
                self.assertEqual(str(run.font.color.rgb), '000000')

    def test_assessed_page_budget(self):
        with pymupdf.open(PDF_FOLDER / f'{STEM}.pdf') as pdf:
            pages = [p.get_text() for p in pdf]
        self.assertIn('Objective and Validation Metrics', pages[1])
        self.assertIn('Analysis and Insights', pages[2])
        self.assertIn('Evaluation of Aims', pages[3])
        self.assertIn('Final Reflection', pages[4])
        self.assertIn('Appendix A', pages[5])
        for number, text in enumerate(pages, 1):
            self.assertEqual(text.split()[0], str(number))
            self.assertGreater(len(text.split()), 8, f'Empty page {number}')

    def test_pdf_uses_black_times_and_a4(self):
        with pymupdf.open(PDF_FOLDER / f'{STEM}.pdf') as pdf:
            for page in pdf:
                self.assertAlmostEqual(page.rect.width, 595.28, delta=0.2)
                self.assertAlmostEqual(page.rect.height, 841.89, delta=0.2)
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines', []):
                        for span in line.get('spans', []):
                            if span['text'].strip():
                                self.assertIn('TimesNewRoman', span['font'])
                                self.assertEqual(span['color'], 0)

    def test_evidence_limits_and_reflection_are_present(self):
        with pymupdf.open(PDF_FOLDER / f'{STEM}.pdf') as pdf:
            text = '\n'.join(p.get_text() for p in pdf)
        compact = re.sub(r'\s+', ' ', text)
        for phrase in ('IP2b', 'Only two interviews were recorded', '15-second',
                       '90%', 'partially confirmed', 'passthrough AR',
                       'Prototype Session Review', 'Methodological Reflection', 'Concept Evaluation',
                       'Improvements and Extensions'):
            self.assertIn(phrase, compact)
        for symbol in (';', '—', '–', '--'):
            self.assertNotIn(symbol, text)

    def test_current_session_and_recording_slots(self):
        with pymupdf.open(PDF_FOLDER / f'{STEM}.pdf') as pdf:
            text = '\n'.join(page.get_text() for page in pdf)
        compact = re.sub(r'\s+', ' ', text)
        for phrase in ('22 October 2026', 'Meta Quest 3S',
                       'XRStudyWhiteboard_Quest_v1.4.apk',
                       'Recording 1 Google Drive link',
                       'Recording 2 Google Drive link'):
            self.assertIn(phrase, compact)
        for participant in range(1, 6):
            self.assertIn(f'P0{participant}: Participant {participant}', compact)
        for previous_name in ('Kai', 'Syifa', 'Serafina', 'Safira'):
            self.assertNotIn(previous_name, text)

    def test_readme_and_wiki_navigation(self):
        readme = (ROOT / 'README.md').read_text()
        self.assertIn('Design%20Evaluation%203%20and%20Reflection.pdf', readme)
        for name in ('Home.md', 'Design-Evaluation-3-and-Reflection.md'):
            self.assertTrue((ROOT / 'wiki' / name).is_file())
        self.assertIn('partially confirmed',
                      (ROOT / 'wiki' / 'Design-Evaluation-3-and-Reflection.md').read_text())


if __name__ == '__main__':
    unittest.main(verbosity=2)
