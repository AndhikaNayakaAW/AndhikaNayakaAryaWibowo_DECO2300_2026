"""Build the APA 7 Evaluation 3 DOCX from the reviewable Markdown source.

Usage: python 'Working Files/Interactive Prototype 2/Tools/build_evaluation3.py'
Export the resulting DOCX to PDF with Word or LibreOffice, then run
check_evaluation3.py. Requires python-docx.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
FOLDER = ROOT / 'Working Files' / 'Interactive Prototype 2' / 'Evaluation'
STEM = 'Design Evaluation 3 and Reflection'
EVIDENCE = ROOT / 'Working Files/Interactive Prototype 2/Earlier Evidence/Iteration2_XRStudyWhiteboard_Transcript_and_Results.md'
FONT = 'Times New Roman'
BLACK = RGBColor(0, 0, 0)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = Inches(.5)


def clean(text):
    return text.replace('—', ',').replace('–', '-').replace('--', '-').replace(';', '.')


def run_format(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold, run.italic = bold, italic
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)


for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.first_line_indent = Inches(0 if name != 'Normal' else .5)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.first_line_indent = Inches(0)
header.paragraph_format.line_spacing = 1
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
field_run = OxmlElement('w:r')
field_text = OxmlElement('w:t')
field_text.text = '1'
field_run.append(field_text)
field.append(field_run)
header._p.append(field)


def inline(paragraph, text, size=12, bold=False, italic=False):
    tokens = re.split(r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*)', clean(text))
    for token in tokens:
        if not token:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', token)
        if link:
            relation = paragraph.part.relate_to(
                link[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True)
            node = OxmlElement('w:hyperlink')
            node.set(qn('r:id'), relation)
            link_run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            for tag, attrs in [('rFonts', {'ascii': FONT, 'hAnsi': FONT}),
                               ('sz', {'val': str(int(size * 2))}),
                               ('color', {'val': '000000'}), ('u', {'val': 'single'})]:
                element = OxmlElement('w:' + tag)
                for key, value in attrs.items():
                    element.set(qn('w:' + key), value)
                properties.append(element)
            link_run.append(properties)
            text_node = OxmlElement('w:t')
            text_node.text = link[1]
            link_run.append(text_node)
            node.append(link_run)
            paragraph._p.append(node)
        else:
            token_bold = token.startswith('**') and token.endswith('**')
            token_italic = not token_bold and token.startswith('*') and token.endswith('*')
            value = token[2:-2] if token_bold else token[1:-1] if token_italic else token
            value = value.replace('`', '')
            run = paragraph.add_run(value)
            run_format(run, size, bold or token_bold, italic or token_italic)


def paragraph(text='', center=False, bold=False, indent=True, page=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(.5 if indent else 0)
    p.paragraph_format.page_break_before = page
    p.paragraph_format.keep_with_next = keep
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    inline(p, text, bold=bold)
    return p


def heading(text, level=1, page=False):
    return paragraph(text, center=level == 1, bold=True, indent=False, page=page, keep=True)


def apa_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available = 6.2677
    widths = ([1.6, available - 1.6] if len(headers) == 2 else
              [.55, 1.0, available - 1.55] if len(headers) == 3 and headers[0] == 'ID' else
              [2, .6, available - 2.6] if len(headers) == 3 else
              [1.6, .5, .6, available - 2.7] if len(headers) == 4 else
              [2.35] + [(available - 2.35) / (len(headers)-1)] * (len(headers)-1))
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn('w:w'), str(round(width * 1440)))
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        element = OxmlElement('w:' + edge)
        element.set(qn('w:val'), 'single' if edge in ('top', 'bottom') else 'nil')
        element.set(qn('w:color'), '000000')
        element.set(qn('w:sz'), '6')
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for row_index, values in enumerate([headers] + rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            props.append(OxmlElement('w:tblHeader'))
        for i, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0)
            inline(p, value, size=10, bold=row_index == 0)
            if row_index == 0:
                cell_borders = OxmlElement('w:tcBorders')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '6')
                bottom.set(qn('w:color'), '000000')
                cell_borders.append(bottom)
                cell._tc.get_or_add_tcPr().append(cell_borders)
    return table


raw_table_counter = 0
participant_count = 0


def render(markdown, raw=False):
    global raw_table_counter, participant_count
    lines = markdown.splitlines()
    i, page_pending = 0, False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line in ('<!-- page -->', '<!-- appendices -->'):
            page_pending = True
            continue
        if line == '<!-- raw-evidence -->':
            heading('Appendix B', page=True)
            heading('Supplied Interview Transcripts and Notes', level=1)
            evidence = EVIDENCE.read_text()
            for participant in range(1, 6):
                marker = f'## P0{participant}'
                start = evidence.index(marker)
                end_marker = f'## P0{participant+1}' if participant < 5 else '## Combined results'
                end = evidence.index(end_marker, start)
                block = evidence[start:end]
                if participant <= 2:
                    title = block.splitlines()[0]
                    transcript = block.split('### Transcript\n', 1)[1]
                    render(title + '\n\n' + transcript, raw=True)
                else:
                    render(block, raw=True)
            continue
        if line.startswith('#'):
            count = len(line) - len(line.lstrip('#'))
            value = line[count:].strip().replace(' — ', ': ')
            if raw:
                is_participant = count == 2 and value.startswith('P0')
                heading(value, level=1 if is_participant else 2,
                        page=is_participant and participant_count > 0)
                if is_participant:
                    participant_count += 1
            else:
                heading(value, level=min(count, 2), page=page_pending)
            page_pending = False
            continue
        if line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            cells = lambda text: [item.strip() for item in text.strip('|').split('|')]
            if raw:
                raw_table_counter += 1
                paragraph(f'**Table B{raw_table_counter}**', indent=False, keep=True)
                paragraph('*Supplied Participant Record*', indent=False, keep=True)
            apa_table(cells(block[0]), [cells(row) for row in block[2:]])
            continue
        table_label = re.fullmatch(r'\*\*Table [A-Z]?\d+\*\*', line)
        table_title = line.startswith('*') and line.endswith('*') and not line.startswith('**')
        if table_label or table_title:
            paragraph(line, indent=False, keep=True)
        elif line.startswith('> '):
            paragraph('*' + line[2:] + '*', indent=False)
        elif line.startswith('- '):
            paragraph(line[2:], indent=False)
        else:
            paragraph(line, indent=not line.startswith('['), page=page_pending)
        page_pending = False


for _ in range(4):
    paragraph('', center=True, indent=False)
paragraph('Design Evaluation 3 and Reflection', center=True, bold=True, indent=False)
paragraph('XR Study Whiteboard', center=True, bold=True, indent=False)
paragraph('Andhika Nayaka Arya Wibowo', center=True, indent=False)
paragraph('The University of Queensland', center=True, indent=False)
paragraph('DECO2300: Digital Prototyping and Extended Reality', center=True, indent=False)
paragraph('DECO2300 Teaching Team', center=True, indent=False)
paragraph('21 October 2026', center=True, indent=False)

source = (FOLDER / f'{STEM}.md').read_text()
render('<!-- page -->\n' + source)
doc.core_properties.title = STEM
doc.core_properties.author = 'Andhika Nayaka Arya Wibowo'
doc.core_properties.subject = 'Final IP2b evaluation with four assessed pages and evidence appendices'
output = FOLDER / f'{STEM}.docx'
doc.save(output)
print(output)
