"""Build the prospective plan, participant form, and retrospective journal."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

PACK = Path(__file__).resolve().parents[1]
FONT = 'Times New Roman'


def draw_sketch():
    image = Image.new('RGB', (1800, 900), 'white')
    draw = ImageDraw.Draw(image)
    font_path = '/System/Library/Fonts/Supplemental/Times New Roman.ttf'
    font = ImageFont.truetype(font_path, 45)
    small = ImageFont.truetype(font_path, 34)
    boxes = [(80,100,760,300), (1040,100,1720,300),
             (80,520,760,760), (1040,520,1720,760)]
    labels = [('LEFT CONTROLLER', 'Move, X to board, Y to tables'),
              ('RIGHT CONTROLLER', 'Select tools and write'),
              ('TEACHING BOARD', 'Shared lesson and action replay'),
              ('FIXED DESK PAPER', 'Personal notes and paper tools')]
    for rect, (title, body) in zip(boxes, labels):
        draw.rectangle(rect, outline='black', width=4)
        cx=(rect[0]+rect[2])/2
        draw.text((cx,rect[1]+55),title,font=font,fill='black',anchor='mm')
        draw.text((cx,rect[1]+125),body,font=small,fill='black',anchor='mm')
    def arrow(a,b):
        draw.line([a,b],fill='black',width=4)
        x,y=b
        draw.polygon([(x,y),(x-12,y-22),(x+12,y-22)],fill='black')
    arrow((420,300),(420,520)); arrow((1380,300),(1380,520))
    draw.line([(760,640),(1040,640)],fill='black',width=4)
    draw.text((900,460),'Look between board and paper',font=font,fill='black',anchor='mm')
    draw.text((900,840),'Retrospective workflow sketch prepared 21 October 2026',font=small,fill='black',anchor='mm')
    image.save(PACK/'Design Journal/Interaction Sketch.png')


def style_run(run, bold=False, italic=False, size=12):
    run.font.name=FONT; run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0)
    run.bold=bold; run.italic=italic


def clean(text):
    text = text.replace('—', ',').replace('–', '-').replace(';', '.').replace('--','-')
    return re.sub(r'_{60,}', '_' * 54, text)


def build(relative):
    source=PACK/relative
    document=Document()
    sec=document.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=Inches(.5)
    normal=document.styles['Normal']
    normal.font.name=FONT; normal.font.size=Pt(12); normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=2
    normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Inches(.5)
    header=sec.header.paragraphs[0]
    header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent=Inches(0)
    header.paragraph_format.line_spacing=1
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text='1'; run.append(text);field.append(run);header._p.append(field)
    def para(text='',bold=False,italic=False,center=False,indent=True,page=False,keep=False):
        p=document.add_paragraph()
        p.paragraph_format.line_spacing=2
        p.paragraph_format.first_line_indent=Inches(.5 if indent else 0)
        p.paragraph_format.page_break_before=page
        p.paragraph_format.keep_with_next=keep
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        style_run(p.add_run(clean(text)),bold,italic)
        return p
    title=source.read_text().splitlines()[0].lstrip('# ')
    for _ in range(4):para('',center=True,indent=False)
    for text,bold in [(title,True),('XR Study Whiteboard',True),('Andhika Nayaka Arya Wibowo',False),
                      ('The University of Queensland',False),('DECO2300: Digital Prototyping and Extended Reality',False),
                      ('DECO2300 Teaching Team',False),('21 October 2026',False)]:
        para(text,bold=bold,center=True,indent=False)
    lines=source.read_text().splitlines(); index=0; page=True
    while index<len(lines):
        line=lines[index].strip();index+=1
        if not line:continue
        if line=='<!-- page -->':page=True;continue
        if line.startswith('<!-- figure:'):
            name=line.removeprefix('<!-- figure:').removesuffix('-->').strip()
            document.add_picture(str(source.parent/name),width=Inches(6))
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            para(line.lstrip('# '),bold=True,center=level==1,indent=False,page=page,keep=True)
            page=False;continue
        if line.startswith('|'):
            rows=[line]
            while index<len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip());index+=1
            split=lambda row:[item.strip() for item in row.strip('|').split('|')]
            headers=split(rows[0]); data=[headers]+[split(r) for r in rows[2:]]
            table=document.add_table(rows=0,cols=len(headers)); table.autofit=False
            widths=[1.55,4.71] if len(headers)==2 else [6.26/len(headers)]*len(headers)
            for column,width in zip(table._tbl.tblGrid.gridCol_lst,widths):column.set(qn('w:w'),str(round(width*1440)))
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','bottom','left','right','insideV','insideH']:
                element=OxmlElement('w:'+edge);element.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil')
                element.set(qn('w:sz'),'6');element.set(qn('w:color'),'000000');borders.append(element)
            table._tbl.tblPr.append(borders)
            for row_index,values in enumerate(data):
                row=table.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if row_index==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell,width,value in zip(row.cells,widths,values):
                    cell.width=Inches(width)
                    p=cell.paragraphs[0];p.paragraph_format.line_spacing=1
                    p.paragraph_format.first_line_indent=Inches(0)
                    style_run(p.add_run(clean(value)),bold=row_index==0,size=10)
                    if row_index==0:
                        borders=OxmlElement('w:tcBorders');bottom=OxmlElement('w:bottom')
                        bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'6');bottom.set(qn('w:color'),'000000')
                        borders.append(bottom);cell._tc.get_or_add_tcPr().append(borders)
            continue
        if line.startswith('**') and line.endswith('**'):
            para(line.strip('*'),bold=True,indent=False,keep=True)
        elif line.startswith('*') and line.endswith('*'):
            para(line.strip('*'),italic=True,indent=False,keep=True)
        else:para(line,indent=not ('____' in line))
    document.core_properties.title=title
    document.core_properties.author='Andhika Nayaka Arya Wibowo'
    document.core_properties.subject='Prospective preparation for the mandatory Week 12 studio'
    output=source.with_suffix('.docx');document.save(output);print(output)


if __name__=='__main__':
    draw_sketch()
    for name in ['Testing/Testing Plan.md','Testing/Participant Record.md','Design Journal/Design Journal.md']:
        build(name)
